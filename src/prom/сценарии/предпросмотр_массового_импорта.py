"""Локальная проверка XLSX с новыми тестовыми товарами без POST."""

from __future__ import annotations

import argparse
import hashlib
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from rich.console import Console
from rich.table import Table

from prom.сценарии.хранилище_снимков import прочитать_снимок

ХАРАКТЕРИСТИКИ = frozenset(
    {
        "Назва_Характеристики",
        "Одиниця_виміру_Характеристики",
        "Значення_Характеристики",
    }
)
ОБЯЗАТЕЛЬНЫЕ = (
    "Код_товару",
    "Назва_позиції",
    "Назва_позиції_укр",
    "Опис",
    "Опис_укр",
    "Ціна",
    "Валюта",
    "Ідентифікатор_товару",
)


@dataclass(frozen=True)
class ИтогПредпросмотра:
    количество: int
    уникальных_артикулов: int
    пересечений: int
    размер: int
    sha256: str
    предупреждения: tuple[str, ...]
    разрешено_отправить: bool = False


def проверить_файл(файл: Path, каталог: Path, ключ: str) -> ИтогПредпросмотра:
    """Проверить ровно 500 новых тестовых товаров по сохранённому снимку."""
    if файл.suffix.lower() != ".xlsx" or not файл.is_file():
        raise ValueError("Нужен существующий XLSX")
    размер = файл.stat().st_size
    if not 0 < размер <= 5_000_000:
        raise ValueError("Недопустимый размер файла")
    снимок, _ = прочитать_снимок(каталог, ключ, "до")
    книга = load_workbook(файл, read_only=True, data_only=True)
    try:
        if len(книга.worksheets) != 1:
            raise ValueError("Нужен один лист")
        строки = книга.worksheets[0].iter_rows(values_only=True)
        заголовки = next(строки, None)
        if заголовки is None or any(
            not isinstance(имя, str) or not имя.strip() for имя in заголовки
        ):
            raise ValueError("Есть пустые заголовки")
        повторения = Counter(заголовки)
        лишние = [
            имя
            for имя, число in повторения.items()
            if число > 1 and имя not in ХАРАКТЕРИСТИКИ
        ]
        if лишние:
            raise ValueError(f"Повторяются обычные заголовки: {лишние}")
        if any(имя not in заголовки for имя in ОБЯЗАТЕЛЬНЫЕ):
            raise ValueError("Отсутствуют обязательные колонки")
        индексы = {имя: заголовки.index(имя) for имя in ОБЯЗАТЕЛЬНЫЕ}
        артикулы: set[str] = set()
        существующие = {
            карточка.артикул for карточка in снимок.карточки if карточка.артикул
        }
        пересечения = 0
        количество = 0
        for номер, значения in enumerate(строки, start=2):
            if not any(значение is not None for значение in значения):
                continue
            количество += 1
            строка = {
                имя: значения[индекс] if индекс < len(значения) else None
                for имя, индекс in индексы.items()
            }
            артикул = строка["Код_товару"]
            if not isinstance(артикул, str) or not артикул.startswith("PROM-LOADTEST-"):
                raise ValueError(f"Строка {номер}: неверный тестовый артикул")
            if артикул in артикулы:
                raise ValueError(f"Строка {номер}: повтор артикула {артикул}")
            артикулы.add(артикул)
            if артикул in существующие:
                пересечения += 1
            if строка["Ідентифікатор_товару"] not in (None, ""):
                raise ValueError(f"Строка {номер}: заполнен ID существующего товара")
            for имя in ("Назва_позиції", "Назва_позиції_укр", "Опис", "Опис_укр"):
                значение = строка[имя]
                if not isinstance(значение, str) or not значение.strip():
                    raise ValueError(f"Строка {номер}: пустое поле {имя}")
            цена = строка["Ціна"]
            if (
                isinstance(цена, bool)
                or not isinstance(цена, (int, float))
                or цена <= 0
            ):
                raise ValueError(f"Строка {номер}: некорректная цена")
            if строка["Валюта"] != "UAH":
                raise ValueError(f"Строка {номер}: валюта должна быть UAH")
        if количество != 500:
            raise ValueError(f"Ожидалось 500 товаров, найдено {количество}")
        if пересечения:
            raise ValueError(f"Артикулы пересекаются со снимком: {пересечения}")
        return ИтогПредпросмотра(
            количество,
            len(артикулы),
            пересечения,
            размер,
            hashlib.sha256(файл.read_bytes()).hexdigest(),
            (
                "Сравнение выполнено только с сохранённым снимком, не с текущим API",
                "Проверены общие поля; принятие всех колонок Prom.ua не гарантировано",
                "Prom.ua может опубликовать импортированные товары",
                "Предпросмотр не отправляет данные и не разрешает POST",
            ),
        )
    finally:
        книга.close()


def показать_предпросмотр(
    итог: ИтогПредпросмотра, *, консоль: Console | None = None
) -> None:
    """Вывести результат в Rich."""
    вывод = консоль if консоль is not None else Console()
    таблица = Table(title="PROM / предпросмотр 500 товаров")
    таблица.add_column("Параметр")
    таблица.add_column("Значение")
    for название, значение in (
        ("Товаров", итог.количество),
        ("Уникальных артикулов", итог.уникальных_артикулов),
        ("Пересечений со снимком", итог.пересечений),
        ("Размер, байт", итог.размер),
        ("SHA-256", итог.sha256),
        ("Отправка в API", "ЗАПРЕЩЕНА"),
    ):
        таблица.add_row(название, str(значение))
    вывод.print(таблица)
    for предупреждение in итог.предупреждения:
        вывод.print(f"[yellow]Внимание:[/yellow] {предупреждение}")


def main() -> None:
    """Команда локального предпросмотра."""
    парсер = argparse.ArgumentParser(description="PROM: предпросмотр 500 товаров")
    парсер.add_argument("файл", type=Path)
    парсер.add_argument("ключ")
    парсер.add_argument("--снимки", type=Path, default=Path("data/снимки_каталога"))
    аргументы = парсер.parse_args()
    показать_предпросмотр(
        проверить_файл(аргументы.файл, аргументы.снимки, аргументы.ключ)
    )


if __name__ == "__main__":
    main()
