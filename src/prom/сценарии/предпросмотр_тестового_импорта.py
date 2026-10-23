"""Предпросмотр одного тестового XLSX без сетевых запросов и изменения магазина."""

from __future__ import annotations

import argparse
import hashlib
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from rich.console import Console
from rich.table import Table

from prom.сценарии.хранилище_снимков import прочитать_снимок


@dataclass(frozen=True)
class Предпросмотр:
    """Локальный результат проверки, не разрешение на импорт."""

    артикул: str
    имя: str
    размер: int
    sha256: str
    предупреждения: tuple[str, ...]
    разрешено_отправить: bool = False


def проверить_файл(файл: Path, каталог: Path, ключ: str) -> Предпросмотр:
    """Проверить один товар и отсутствие артикула в снимке «до»."""
    if файл.suffix.lower() != ".xlsx" or not файл.is_file():
        raise ValueError("Нужен существующий файл XLSX")
    размер = файл.stat().st_size
    if размер == 0 or размер > 5_000_000:
        raise ValueError("Пустой или слишком большой тестовый файл")
    снимок, _ = прочитать_снимок(каталог, ключ, "до")
    книга = load_workbook(файл, read_only=True, data_only=True)
    try:
        if len(книга.worksheets) != 1:
            raise ValueError("Тестовый файл должен содержать один лист")
        строки = книга.worksheets[0].iter_rows(values_only=True)
        заголовки = next(строки, None)
        if заголовки is None or len(заголовки) != len(set(заголовки)):
            raise ValueError("Заголовки отсутствуют или повторяются")
        требуемые = ("Код_товару", "Назва_позиції", "Ціна", "Валюта")
        if not all(название in заголовки for название in требуемые):
            raise ValueError("В файле отсутствуют обязательные колонки")
        данные = [строка for строка in строки if any(значение is not None for значение in строка)]
        if len(данные) != 1:
            raise ValueError("Тестовый файл должен содержать ровно один товар")
        строка = dict(zip(заголовки, данные[0], strict=False))
        артикул = строка["Код_товару"]
        имя = строка["Назва_позиції"]
        if not isinstance(артикул, str) or not артикул.startswith("PROM-TEST-"):
            raise ValueError("Артикул должен начинаться с PROM-TEST-")
        if not isinstance(имя, str) or not имя.strip():
            raise ValueError("Название товара обязательно")
        if строка["Ціна"] != 1 or строка["Валюта"] != "UAH":
            raise ValueError("Тестовая цена должна быть 1 UAH")
        if any(карточка.артикул == артикул for карточка in снимок.карточки):
            raise ValueError("Артикул уже присутствует в снимке каталога")
        if строка.get("Ідентифікатор_товару") not in (None, ""):
            raise ValueError("Новый товар не должен содержать ID существующей карточки")
        предупреждения = (
            "Уникальность проверена только по сохранённому снимку, не по текущему API",
            "Prom.ua может опубликовать товар при импорте независимо от признака публикации",
            "POST и изменение товаров заблокированы",
        )
        сумма = hashlib.sha256(файл.read_bytes()).hexdigest()
        return Предпросмотр(артикул, имя, размер, сумма, предупреждения)
    finally:
        книга.close()


def показать_предпросмотр(итог: Предпросмотр, *, консоль: Console | None = None) -> None:
    """Rich-отчёт с контрольной суммой для последующего подтверждения."""
    вывод = консоль if консоль is not None else Console()
    таблица = Table(title="PROM / предварительный просмотр тестового импорта")
    таблица.add_column("Параметр")
    таблица.add_column("Значение")
    таблица.add_row("Артикул", итог.артикул)
    таблица.add_row("Название", итог.имя)
    таблица.add_row("Размер, байт", str(итог.размер))
    таблица.add_row("SHA-256", итог.sha256)
    таблица.add_row("Отправка в API", "ЗАПРЕЩЕНА")
    вывод.print(таблица)
    for предупреждение in итог.предупреждения:
        вывод.print(f"[yellow]Внимание:[/yellow] {предупреждение}")


def main() -> None:
    """Проверить файл без обращения к Prom.ua."""
    парсер = argparse.ArgumentParser(description="PROM: локальный предпросмотр XLSX")
    парсер.add_argument("файл", type=Path)
    парсер.add_argument("ключ", help="Ключ ранее созданного снимка 'до'")
    парсер.add_argument("--снимки", type=Path, default=Path("data/снимки_каталога"))
    аргументы = парсер.parse_args()
    показать_предпросмотр(проверить_файл(аргументы.файл, аргументы.снимки, аргументы.ключ))


if __name__ == "__main__":
    main()
