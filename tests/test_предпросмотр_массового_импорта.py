"""Проверки безопасного предпросмотра небольших партий."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from prom.сценарии.предпросмотр_массового_импорта import проверить_файл
from prom.сценарии.сверка_каталога import Карточка, Снимок
from prom.сценарии.хранилище_снимков import сохранить_снимок


@pytest.fixture
def подготовка(tmp_path: Path) -> tuple[Path, Path]:
    """Создать две новые карточки и полный снимок для теста."""
    сохранить_снимок(
        tmp_path, "test", "до", Снимок((Карточка(1, "OTHER", None, False),), True)
    )
    книга = Workbook()
    лист = книга.active
    assert лист is not None
    лист.append([
        "Код_товару", "Назва_позиції", "Назва_позиції_укр", "Опис",
        "Опис_укр", "Ціна", "Валюта", "Ідентифікатор_товару",
        "Оптова_ціна", "Мінімальне_замовлення_опт",
    ])
    for номер in (1, 2):
        лист.append([
            f"PROM-LOADTEST-{номер:04d}", "Название", "Назва", "Описание",
            "Опис", 2.0, "UAH", None, 1.5, 2,
        ])
    путь = tmp_path / "товары.xlsx"
    книга.save(путь)
    return путь, tmp_path


def test_две_карточки_разрешены_только_для_предпросмотра(
    подготовка: tuple[Path, Path],
) -> None:
    файл, каталог = подготовка
    итог = проверить_файл(файл, каталог, "test", количество_ожидаемое=2)
    assert итог.количество == 2
    assert итог.уникальных_артикулов == 2
    assert not итог.разрешено_отправить


def test_по_умолчанию_ожидаются_500(
    подготовка: tuple[Path, Path],
) -> None:
    файл, каталог = подготовка
    with pytest.raises(ValueError, match="Ожидалось 500"):
        проверить_файл(файл, каталог, "test")


@pytest.mark.parametrize("число", [0, -1, 501])
def test_недопустимое_ожидаемое_количество(
    подготовка: tuple[Path, Path], число: int,
) -> None:
    файл, каталог = подготовка
    with pytest.raises(ValueError, match="от 1 до 500"):
        проверить_файл(файл, каталог, "test", количество_ожидаемое=число)


@pytest.mark.parametrize("значение", [None, 0, -1, True, "1.5", 3])
def test_некорректная_оптовая_цена(
    подготовка: tuple[Path, Path], значение: object,
) -> None:
    файл, каталог = подготовка
    from openpyxl import load_workbook

    книга = load_workbook(файл)
    лист = книга.active
    assert лист is not None
    лист.cell(2, 9).value = значение
    книга.save(файл)
    with pytest.raises(ValueError, match="оптовая цена"):
        проверить_файл(файл, каталог, "test", количество_ожидаемое=2)


@pytest.mark.parametrize("значение", [None, 0, 1, -1, True, 2.5, "2"])
def test_некорректный_оптовый_минимум(
    подготовка: tuple[Path, Path], значение: object,
) -> None:
    файл, каталог = подготовка
    from openpyxl import load_workbook

    книга = load_workbook(файл)
    лист = книга.active
    assert лист is not None
    лист.cell(2, 10).value = значение
    книга.save(файл)
    with pytest.raises(ValueError, match="минимум оптового заказа"):
        проверить_файл(файл, каталог, "test", количество_ожидаемое=2)
